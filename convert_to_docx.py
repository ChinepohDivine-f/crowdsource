"""
Script to convert the Senzor System Report from Markdown to Word (.docx)
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re

def add_heading(doc, text, level=1):
    """Add a styled heading"""
    heading = doc.add_heading(text, level=level)
    return heading

def add_paragraph(doc, text, style='Normal'):
    """Add a paragraph with optional style"""
    p = doc.add_paragraph(text, style=style)
    return p

def add_code_block(doc, code):
    """Add a code block with monospace font"""
    p = doc.add_paragraph()
    run = p.add_run(code)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.5)
    # Shading logic removed to prevent AttributeError
    # shading = p._element.get_or_add_pPr()
    # shading_elm = shading.get_or_add_shd()
    # shading_elm.set(qn('w:fill'), 'F5F5F5')
    return p

def add_table(doc, headers, rows):
    """Add a table"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light List Accent 1'
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        # Bold header
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    # Data rows
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = str(cell_data)
    
    return table

def parse_markdown_to_docx(md_file, docx_file):
    """Convert Markdown to Word document"""
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_buffer = []
    in_table = False
    table_headers = []
    table_rows = []
    
    while i < len(lines):
        line = lines[i]
        
        # Code blocks
        if line.startswith('```'):
            if in_code_block:
                # End of code block
                add_code_block(doc, '\n'.join(code_buffer))
                code_buffer = []
                in_code_block = False
            else:
                # Start of code block
                in_code_block = True
            i += 1
            continue
        
        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue
        
        # Tables
        if line.startswith('|') and '|' in line[1:]:
            cells = [cell.strip() for cell in line.split('|')[1:-1]]
            
            if not in_table:
                # First row (headers)
                table_headers = cells
                in_table = True
                i += 1
                # Skip separator line
                if i < len(lines) and lines[i].startswith('|--'):
                    i += 1
                continue
            else:
                # Data row
                table_rows.append(cells)
                i += 1
                # Check if next line is still table
                if i >= len(lines) or not lines[i].startswith('|'):
                    # End of table
                    add_table(doc, table_headers, table_rows)
                    doc.add_paragraph()  # Spacing
                    in_table = False
                    table_headers = []
                    table_rows = []
                continue
        
        # Headings
        if line.startswith('# '):
            add_heading(doc, line[2:], level=1)
        elif line.startswith('## '):
            add_heading(doc, line[3:], level=2)
        elif line.startswith('### '):
            add_heading(doc, line[4:], level=3)
        elif line.startswith('#### '):
            add_heading(doc, line[5:], level=4)
        
        # Horizontal rules
        elif line.strip() == '---':
            doc.add_paragraph()  # Just add spacing
        
        # Bullet lists
        elif line.startswith('- ') or line.startswith('* '):
            text = line[2:].strip()
            # Remove markdown formatting
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)  # Bold
            text = re.sub(r'\*(.*?)\*', r'\1', text)  # Italic
            text = re.sub(r'`(.*?)`', r'\1', text)  # Code
            p = doc.add_paragraph(text, style='List Bullet')
        
        # Numbered lists
        elif re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line).strip()
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'\*(.*?)\*', r'\1', text)
            text = re.sub(r'`(.*?)`', r'\1', text)
            p = doc.add_paragraph(text, style='List Number')
        
        # Regular paragraphs
        elif line.strip() and not line.startswith(('#', '```', '|', '[')):
            text = line.strip()
            # Remove markdown formatting
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'\*(.*?)\*', r'\1', text)
            text = re.sub(r'`(.*?)`', r'\1', text)
            
            if text:
                p = doc.add_paragraph(text)
        
        i += 1
    
    # Save document
    doc.save(docx_file)
    print(f"✅ Word document created: {docx_file}")

if __name__ == '__main__':
    md_file = 'SYSTEM_REPORT_FULL.md'
    docx_file = 'Senzor_System_Report_Full.docx'
    
    parse_markdown_to_docx(md_file, docx_file)
